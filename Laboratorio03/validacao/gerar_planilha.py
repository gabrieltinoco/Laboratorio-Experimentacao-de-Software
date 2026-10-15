"""Sorteio da amostra-ouro e planilhas de rotulacao, uma por avaliador.

Uso: python -m validacao.gerar_planilha --config config.yaml

As planilhas nao trazem a saida da heuristica de release corretiva, para nao
influenciar os avaliadores. Planilhas ja existentes nao sao sobrescritas.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from urllib.parse import quote

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from validacao.rotulos import ABA_RELEASES, ABA_REPOSITORIOS, DIMENSOES, caminho_rotulos

COLUNAS_REPOSITORIOS = ["repositorio", "link", "tipo_projeto", "entregas_reais", "observacoes"]
COLUNAS_RELEASES = ["repositorio", "tag_name", "publicada_em", "link_release", "link_commits",
                    "corretiva", "observacoes"]

INSTRUCOES = [
    "Preencha sem consultar o outro avaliador e sem olhar a saida da heuristica.",
    "",
    "Aba repositorios (uma linha por repositorio):",
    "  tipo_projeto: biblioteca/framework | aplicação/serviço | ferramenta CLI | outro",
    "  entregas_reais: as releases representam entregas reais ao usuario? sim | não | incerto",
    "",
    "Aba releases (5 releases sorteadas por repositorio):",
    "  corretiva: o objetivo principal da release e corrigir um defeito recente? sim | não",
    "  Leia as release notes (link_release) e a lista de commits (link_commits).",
    "",
    "Use a coluna observacoes para justificar casos dificeis. Criterios em validacao/PROTOCOLO.md.",
]


def ler_releases(caminho: str | Path) -> pd.DataFrame:
    with Path(caminho).open(encoding="utf-8") as arquivo:
        return pd.DataFrame([json.loads(linha) for linha in arquivo if linha.strip()])


def sortear_repositorios(repositorios: pd.DataFrame, n: int, semente: int) -> list[str]:
    ordenados = repositorios.sort_values("repositorio", ignore_index=True)
    return ordenados.sample(n=n, random_state=semente)["repositorio"].tolist()


def sortear_releases(releases: pd.DataFrame, repositorio: str, k: int, semente: int,
                     inicio: str, fim: str) -> pd.DataFrame:
    """Sorteia k releases principais da janela e anota a release anterior de cada uma."""
    principais = releases[
        (releases["repositorio"] == repositorio)
        & ~releases["draft"].astype(bool)
        & ~releases["prerelease"].astype(bool)
        & releases["published_at"].notna()
    ].sort_values(["published_at", "tag_name"], ignore_index=True)
    principais["tag_anterior"] = principais["tag_name"].shift(1)
    data = principais["published_at"].str[:10]
    na_janela = principais[(data >= inicio) & (data <= fim)]
    sorteadas = na_janela.sample(n=min(k, len(na_janela)), random_state=semente)
    return sorteadas.sort_values("published_at", ignore_index=True)


def link_commits(repositorio: str, anterior: str | None, tag: str) -> str:
    if not isinstance(anterior, str):
        return f"https://github.com/{repositorio}/commits/{quote(tag, safe='')}"
    return f"https://github.com/{repositorio}/compare/{quote(anterior, safe='')}...{quote(tag, safe='')}"


def montar_amostra(repositorios: pd.DataFrame, releases: pd.DataFrame, config: dict) -> pd.DataFrame:
    """Uma linha por release sorteada, com os links para rotulacao."""
    parametros = config["validacao"]
    inicio, fim = str(config["janela"]["inicio"]), str(config["janela"]["fim"])
    linhas = []
    for repositorio in sortear_repositorios(repositorios, parametros["repositorios"], parametros["semente"]):
        sorteadas = sortear_releases(releases, repositorio, parametros["releases_por_repositorio"],
                                     parametros["semente"], inicio, fim)
        for _, release in sorteadas.iterrows():
            linhas.append({
                "repositorio": repositorio,
                "tag_name": release["tag_name"],
                "publicada_em": release["published_at"][:10],
                "tag_anterior": release["tag_anterior"] if isinstance(release["tag_anterior"], str) else "",
                "link_release": release["html_url"],
                "link_commits": link_commits(repositorio, release["tag_anterior"], release["tag_name"]),
            })
    return pd.DataFrame(linhas)


def _aba(planilha, titulo: str, colunas: list[str], linhas: list[dict]):
    aba = planilha.create_sheet(titulo)
    aba.append(colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        aba.append([linha.get(coluna, "") for coluna in colunas])
    for indice, coluna in enumerate(colunas, start=1):
        if coluna.startswith("link"):
            for (celula,) in aba.iter_rows(min_row=2, min_col=indice, max_col=indice):
                celula.hyperlink = celula.value
                celula.style = "Hyperlink"
        letra = aba.cell(row=1, column=indice).column_letter
        aba.column_dimensions[letra].width = 60 if coluna.startswith("link") else 24
    aba.freeze_panes = "A2"
    return aba


def _lista_suspensa(aba, colunas: list[str], coluna: str, valores: tuple[str, ...]) -> None:
    validacao = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
    letra = aba.cell(row=1, column=colunas.index(coluna) + 1).column_letter
    validacao.add(f"{letra}2:{letra}{max(aba.max_row, 2)}")
    aba.add_data_validation(validacao)


def escrever_planilha(amostra: pd.DataFrame, caminho: str | Path) -> None:
    planilha = Workbook()
    instrucoes = planilha.active
    instrucoes.title = "instrucoes"
    for texto in INSTRUCOES:
        instrucoes.append([texto])
    instrucoes.column_dimensions["A"].width = 100

    repositorios = [
        {"repositorio": nome, "link": f"https://github.com/{nome}"}
        for nome in amostra["repositorio"].drop_duplicates()
    ]
    abas = {
        ABA_REPOSITORIOS: (_aba(planilha, ABA_REPOSITORIOS, COLUNAS_REPOSITORIOS, repositorios),
                           COLUNAS_REPOSITORIOS),
        ABA_RELEASES: (_aba(planilha, ABA_RELEASES, COLUNAS_RELEASES, amostra.to_dict("records")),
                       COLUNAS_RELEASES),
    }
    for nome_aba, coluna, valores in DIMENSOES.values():
        aba, colunas = abas[nome_aba]
        _lista_suspensa(aba, colunas, coluna, valores)

    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    planilha.save(caminho)


def gerar(config: dict, sobrescrever: bool = False) -> pd.DataFrame:
    dados = Path(config["caminhos"]["dados"])
    pasta = Path(config["validacao"]["pasta"])
    amostra = montar_amostra(pd.read_csv(dados / "repositorios.csv"),
                             ler_releases(dados / "releases.jsonl"), config)

    pasta.mkdir(parents=True, exist_ok=True)
    amostra.to_csv(pasta / "amostra_ouro.csv", index=False)
    for avaliador in config["validacao"]["avaliadores"]:
        caminho = caminho_rotulos(pasta, avaliador)
        if caminho.exists() and not sobrescrever:
            print(f"Mantida (ja existe): {caminho}")
            continue
        escrever_planilha(amostra, caminho)
        print(f"Gerada: {caminho}")
    return amostra


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Sorteia a amostra-ouro e gera as planilhas.")
    parser.add_argument("--config", default="config.yaml")
    parser.add_argument("--sobrescrever", action="store_true",
                        help="regera as planilhas existentes, apagando rotulos ja preenchidos")
    args = parser.parse_args(argv)
    config = yaml.safe_load(Path(args.config).read_text(encoding="utf-8"))
    amostra = gerar(config, args.sobrescrever)
    print(f"{amostra['repositorio'].nunique()} repositorios, {len(amostra)} releases sorteadas.")


if __name__ == "__main__":
    main()
