"""Testes do sorteio da amostra-ouro, da concordancia e do consenso."""

from __future__ import annotations

import json
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from validacao import concordancia, consenso, gerar_planilha
from validacao.rotulos import caminho_rotulos, ler_rotulos, problemas_de_preenchimento


def release(tag, publicada, prerelease=False, draft=False):
    return {"tag_name": tag, "published_at": publicada, "prerelease": prerelease, "draft": draft,
            "html_url": f"https://github.com/o/r/releases/tag/{tag}"}


@pytest.fixture
def config(tmp_path):
    return {
        "janela": {"inicio": "2025-10-01", "fim": "2026-09-30"},
        "caminhos": {"dados": str(tmp_path / "dados")},
        "validacao": {"semente": 42, "repositorios": 3, "releases_por_repositorio": 2,
                      "avaliadores": ["av1", "av2"], "pasta": str(tmp_path / "validacao")},
    }


@pytest.fixture
def dados(config, tmp_path):
    pasta = tmp_path / "dados"
    pasta.mkdir()
    nomes = [f"o/r{i}" for i in range(5)]
    pd.DataFrame({"repositorio": nomes}).to_csv(pasta / "repositorios.csv", index=False)
    with (pasta / "releases.jsonl").open("w", encoding="utf-8") as arquivo:
        for nome in nomes:
            linhas = [
                release("v0.9", "2025-09-01T00:00:00Z"),
                release("v1.0", "2025-11-01T00:00:00Z"),
                release("v1.1-rc", "2025-11-20T00:00:00Z", prerelease=True),
                release("v1.1", "2025-12-01T00:00:00Z"),
                release("v1.2", "2026-01-01T00:00:00Z"),
            ]
            for linha in linhas:
                arquivo.write(json.dumps({"repositorio": nome, **linha}) + "\n")
    return pasta


def preencher(caminho, tipos, entregas, corretivas):
    planilha = load_workbook(caminho)
    for aba, coluna, valores in (("repositorios", 3, tipos), ("repositorios", 4, entregas),
                                 ("releases", 6, corretivas)):
        for linha, valor in enumerate(valores, start=2):
            planilha[aba].cell(row=linha, column=coluna, value=valor)
    planilha.save(caminho)


def test_sorteio_e_reproduzivel_e_ignora_pre_releases_e_releases_fora_da_janela(config, dados):
    primeira = gerar_planilha.gerar(config)
    segunda = gerar_planilha.montar_amostra(pd.read_csv(dados / "repositorios.csv"),
                                            gerar_planilha.ler_releases(dados / "releases.jsonl"), config)

    pd.testing.assert_frame_equal(primeira, segunda)
    assert primeira["repositorio"].nunique() == 3
    assert primeira.groupby("repositorio").size().eq(2).all()
    assert set(primeira["tag_name"]) <= {"v1.0", "v1.1", "v1.2"}


def test_link_de_commits_compara_com_a_release_anterior_mesmo_fora_da_janela(config, dados):
    amostra = gerar_planilha.gerar(config)
    v10 = amostra[amostra["tag_name"] == "v1.0"]
    v11 = amostra[amostra["tag_name"] == "v1.1"]

    assert (v10["link_commits"].str.endswith("/compare/v0.9...v1.0")).all()
    assert (v11["tag_anterior"] == "v1.0").all()
    assert gerar_planilha.link_commits("o/r", None, "v1") == "https://github.com/o/r/commits/v1"


def test_planilha_tem_links_listas_e_nao_e_sobrescrita(config, dados):
    gerar_planilha.gerar(config)
    caminho = caminho_rotulos(config["validacao"]["pasta"], "av1")
    planilha = load_workbook(caminho)

    assert planilha["releases"]["D2"].hyperlink is not None
    assert len(planilha["releases"].data_validations.dataValidation) == 1
    assert len(planilha["repositorios"].data_validations.dataValidation) == 2

    preencher(caminho, ["outro"] * 3, ["sim"] * 3, ["não"] * 6)
    gerar_planilha.gerar(config)
    assert ler_rotulos(caminho)["repositorios"]["tipo_projeto"].eq("outro").all()


def test_problemas_de_preenchimento_apontam_vazios_e_valores_invalidos(config, dados):
    gerar_planilha.gerar(config)
    caminho = caminho_rotulos(config["validacao"]["pasta"], "av1")
    preencher(caminho, ["outro", "lib", None], ["sim"] * 3, ["não"] * 6)

    problemas = problemas_de_preenchimento(ler_rotulos(caminho))
    assert len(problemas) == 2
    assert any("'lib'" in p for p in problemas)


@pytest.mark.parametrize("kappa, esperado", [
    (-0.1, "pobre"), (0.0, "leve"), (0.2, "leve"), (0.35, "razoável"), (0.5, "moderada"),
    (0.75, "substancial"), (0.9, "quase perfeita"), (float("nan"), "indefinido"),
])
def test_interpretacao_de_landis_e_koch(kappa, esperado):
    assert concordancia.landis_koch(kappa) == esperado


def test_kappa_de_cohen_e_de_fleiss_com_valores_calculados_a_mao():
    # Concordancia observada 0,8; esperada 0,5 -> kappa de Cohen = 0,6.
    a = pd.Series(["sim"] * 5 + ["não"] * 5)
    b = pd.Series(["sim"] * 4 + ["não"] * 5 + ["sim"])
    assert concordancia.kappa_cohen(a, b) == pytest.approx(0.6)
    assert concordancia.kappa_fleiss(pd.DataFrame({"a": a, "b": b})) == pytest.approx(0.6)


def test_kappa_indefinido_quando_todos_usam_o_mesmo_rotulo():
    iguais = pd.Series(["sim"] * 4)
    assert math.isnan(concordancia.kappa_cohen(iguais, iguais))
    assert math.isnan(concordancia.kappa_fleiss(pd.DataFrame({"a": iguais, "b": iguais})))


@pytest.fixture
def rotulados(config, dados):
    gerar_planilha.gerar(config)
    pasta = config["validacao"]["pasta"]
    preencher(caminho_rotulos(pasta, "av1"), ["outro", "ferramenta CLI", "outro"], ["sim"] * 3,
              ["sim", "não", "não", "não", "sim", "não"])
    preencher(caminho_rotulos(pasta, "av2"), ["outro", "outro", "outro"], ["sim"] * 3,
              ["sim", "não", "sim", "não", "sim", "não"])
    return concordancia.carregar_rotulos(config)


def test_comparar_resume_kappas_e_lista_divergencias(rotulados):
    resumo, divergencias = concordancia.comparar(rotulados)
    resumo = resumo.set_index("dimensao")

    assert resumo.loc["corretiva", "itens"] == 6
    assert resumo.loc["corretiva", "concordancia_observada"] == pytest.approx(5 / 6, abs=1e-4)
    assert resumo.loc["entregas_reais", "interpretacao"] == "indefinido"
    assert len(divergencias) == 2
    assert set(divergencias["dimensao"]) == {"tipo_projeto", "corretiva"}
    assert divergencias.loc[divergencias["dimensao"] == "tipo_projeto", "tag_name"].eq("").all()


def test_carregar_rotulos_rejeita_planilha_incompleta(config, dados):
    gerar_planilha.gerar(config)
    with pytest.raises(ValueError, match="Planilhas incompletas"):
        concordancia.carregar_rotulos(config)


def test_preservar_decisoes_mantem_o_que_foi_decidido_na_reuniao(rotulados):
    _, divergencias = concordancia.comparar(rotulados)
    anteriores = divergencias.copy()
    anteriores.loc[0, ["rotulo_consenso", "motivo"]] = ["outro", "README descreve um app"]

    preservadas = concordancia.preservar_decisoes(divergencias, anteriores)
    assert list(preservadas.columns) == list(divergencias.columns)
    assert preservadas.loc[0, "rotulo_consenso"] == "outro"
    assert preservadas.loc[1, "rotulo_consenso"] == ""


def test_consenso_usa_rotulo_comum_e_decisoes_da_reuniao(rotulados):
    _, divergencias = concordancia.comparar(rotulados)
    divergencias["rotulo_consenso"] = ["ferramenta CLI", "sim"]
    divergencias["motivo"] = ["instalada via pip como comando", "release notes citam hotfix"]

    final = consenso.consolidar(rotulados, divergencias)
    assert list(final["repositorios"].columns) == ["repositorio", "tipo_projeto", "entregas_reais"]
    assert final["repositorios"]["tipo_projeto"].value_counts()["ferramenta CLI"] == 1
    assert final["releases"]["corretiva"].tolist().count("sim") == 3


def test_consenso_exige_decisao_valida_e_motivo(rotulados):
    _, divergencias = concordancia.comparar(rotulados)
    divergencias["rotulo_consenso"] = ["ferramenta CLI", "talvez"]
    divergencias["motivo"] = ["", "sem acordo"]

    with pytest.raises(ValueError) as erro:
        consenso.consolidar(rotulados, divergencias)
    assert "tipo_projeto" in str(erro.value) and "corretiva" in str(erro.value)
